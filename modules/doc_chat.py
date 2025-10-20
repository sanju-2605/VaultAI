from langchain.docstore.document import Document
from langchain.text_splitter import CharacterTextSplitter
from langchain_community.embeddings import SentenceTransformerEmbeddings
from langchain_community.vectorstores import Chroma

from transformers import pipeline
from docx import Document as DocxDocument
import fitz  # PyMuPDF

embeddings = SentenceTransformerEmbeddings(model_name="all-MiniLM-L6-v2")

def load_docx(file_path):
    doc = DocxDocument(file_path)
    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    return text

def load_pdf(file_path):
    pdf = fitz.open(file_path)
    pages = [pdf[i].get_text() for i in range(len(pdf))]
    text = "\n".join(pages)
    return text

def answer_question_langchain(file_path, file_type, question, k=3):
    # 1. Load document
    if file_type == "docx":
        text = load_docx(file_path)
    elif file_type == "pdf":
        text = load_pdf(file_path)
    else:
        raise ValueError("Unsupported format for LangChain QA")
    # 2. Split text for retrieval
    splitter = CharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
    chunks = splitter.split_text(text)
    docs = [Document(page_content=chunk) for chunk in chunks]
    # 3. Vectorstore & retrieval (semantic search for best chunks)
    db = Chroma.from_documents(docs, embeddings)
    docs_relevant = db.similarity_search(question, k=k)
    # 4. Run QA pipeline on top chunks and pick best answer
    hf_qa_pipe = pipeline("question-answering", model="distilbert-base-cased-distilled-squad")
    best = {"answer": "", "score": 0, "context": ""}
    for doc in docs_relevant:
        result = hf_qa_pipe(question=question, context=doc.page_content)
        if result["score"] > best["score"]:
            best = {
                "answer": result["answer"],
                "score": result["score"],
                "context": doc.page_content[:300]
            }
    if not best["answer"].strip():
        return "No answer found.", ""
    return best["answer"], best["context"]
